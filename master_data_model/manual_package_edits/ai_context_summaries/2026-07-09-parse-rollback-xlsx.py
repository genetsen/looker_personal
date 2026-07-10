"""Parse the exported rollback Manual Data Editor workbook.

This reads the local XLSX exported from Google Drive and writes compact CSV
evidence files for rows that look manually touched plus rows matching the
Purely Elizabeth / Columbus Circle DOOH package. It does not call Google APIs
or mutate any workbook.
"""

from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(
    "/Users/eugenetsenter/Looker_clonedRepo/looker_personal/master_data_model/manual_package_edits"
)
OUT_DIR = ROOT / "ai_context_summaries"
XLSX_PATH = OUT_DIR / "2026-07-09-rollback-copy-export.xlsx"


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def dedupe_headers(raw_headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    headers: list[str] = []
    for index, header in enumerate(raw_headers, start=1):
        name = header.strip() if header else f"blank_header_{index}"
        if name in seen:
            seen[name] += 1
            name = f"{name}__{seen[name]}"
        else:
            seen[name] = 1
        headers.append(name)
    return headers


def safe(row: dict[str, str], name: str) -> str:
    return row.get(name, "").strip()


def write_csv(path: Path, rows: list[dict[str, str]], fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    workbook = openpyxl.load_workbook(XLSX_PATH, data_only=True, read_only=False)
    tab_path = OUT_DIR / "2026-07-09-rollback-copy-xlsx-tabs.csv"
    with tab_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["tab_name", "max_row", "max_column"])
        for sheet in workbook.worksheets:
            writer.writerow([sheet.title, sheet.max_row, sheet.max_column])

    sheet = workbook["Package Editor"]
    raw_headers = [clean(cell.value) for cell in sheet[4]]
    headers = dedupe_headers(raw_headers)
    header_path = OUT_DIR / "2026-07-09-rollback-copy-xlsx-headers.csv"
    with header_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["column_index", "header"])
        for index, header in enumerate(headers, start=1):
            writer.writerow([index, header])

    rows: list[dict[str, str]] = []
    for sheet_row, values in enumerate(
        sheet.iter_rows(min_row=5, max_row=sheet.max_row, values_only=True),
        start=5,
    ):
        row = {"sheet_row": str(sheet_row)}
        row.update({header: clean(value) for header, value in zip(headers, values)})
        rows.append(row)

    marker_headers = [
        header
        for header in headers
        if (
            "manual marker" in header.lower()
            or "manual_marker" in header.lower()
            or header.lower().startswith("man_")
        )
        and header
        not in {
            "Manually Edited?",
            "Manual Edit At",
            "Manual Edit By",
            "Manual Edit Published At",
        }
    ]

    compact_headers = [
        "sheet_row",
        "Package ID",
        "Advertiser",
        "Site",
        "Package Friendly Name",
        "Package Name",
        "Campaign Name",
        "Flight Start Date",
        "Flight End Date",
        "Planned Spend",
        "Planned Impressions",
        "Spend",
        "Impressions",
        "Clicks",
        "Video Plays",
        "Video Completions",
        "Delivery Override Start Date",
        "Delivery Override End Date",
        "Benchmark KPI",
        "Benchmark Value",
        "Manually Edited?",
        "Manual Edit At",
        "Manual Edit By",
        "Manual Edit Published At",
        "Validation Status",
        "Validation Reason",
    ]
    compact_headers.extend(marker_headers)
    compact_headers = [header for header in compact_headers if header == "sheet_row" or header in headers]

    manual_rows: list[dict[str, str]] = []
    key_rows: list[dict[str, str]] = []
    for row in rows:
        package_id = safe(row, "Package ID")
        if not package_id:
            continue

        status = safe(row, "Validation Status").lower()
        manually_edited = safe(row, "Manually Edited?").lower()
        marker_true = any(
            safe(row, marker).lower() in {"true", "yes", "1"} for marker in marker_headers
        )
        has_manual_stamp = any(
            safe(row, name)
            for name in ["Manual Edit At", "Manual Edit By", "Manual Edit Published At"]
        )
        looks_manual = (
            status in {"valid", "blocked"}
            or manually_edited in {"yes", "blocked", "true"}
            or marker_true
            or has_manual_stamp
        )

        searchable = " ".join(
            [
                package_id,
                safe(row, "Package Friendly Name"),
                safe(row, "Package Name"),
                safe(row, "Campaign Name"),
                safe(row, "Advertiser"),
                safe(row, "Site"),
            ]
        ).lower()
        key_match = (
            "ccdooh" in searchable
            or "columbus circle" in searchable
            or "purelyelizabeth" in searchable
            or "purely elizabeth" in searchable
        )

        compact = {header: row.get(header, "") for header in compact_headers}
        if looks_manual:
            manual_rows.append(compact)
        if key_match:
            key_rows.append(compact)

    write_csv(
        OUT_DIR / "2026-07-09-rollback-copy-xlsx-manual-like-rows.csv",
        manual_rows,
        compact_headers,
    )
    write_csv(
        OUT_DIR / "2026-07-09-rollback-copy-xlsx-ccdooh-columbus-rows.csv",
        key_rows,
        compact_headers,
    )

    print(
        "rollback_xlsx",
        f"rows={len(rows)}",
        f"cols={len(headers)}",
        f"manual_like={len(manual_rows)}",
        f"key_rows={len(key_rows)}",
        f"marker_cols={len(marker_headers)}",
    )


if __name__ == "__main__":
    main()
